from pptx import Presentation
from pptx.util import Inches

# Create a presentation object
prs = Presentation()

# Slide 1: Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
title.text = "Project Management Office (PMO) Activities"
subtitle = slide.placeholders[1]
subtitle.text = "Impact on Project Success"

# Slide 2: Introduction
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Introduction"
slide.placeholders[1].text = (
    "The Project Management Office (PMO) plays a crucial role in ensuring project success. "
    "It standardizes processes, provides governance, and aligns projects with business objectives."")

# Slide 3: Key PMO Activities
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Key PMO Activities"
content = (
    "1. Governance & Compliance\n"
    "2. Risk Management\n"
    "3. Resource Allocation\n"
    "4. Project Performance Monitoring\n"
    "5. Process Standardization\n"
    "6. Training & Knowledge Sharing\n"
    "7. Portfolio Management"
)
slide.placeholders[1].text = content

# Slide 4: Impact on Project Success
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Impact on Project Success"
content = (
    "✔ Improved decision-making through structured governance\n"
    "✔ Better risk identification and mitigation\n"
    "✔ Efficient use of resources\n"
    "✔ Standardized processes ensure consistency\n"
    "✔ Enhanced communication and collaboration\n"
    "✔ Increased project success rates"
)
slide.placeholders[1].text = content

# Slide 5: Conclusion
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Conclusion"
slide.placeholders[1].text = (
    "A well-structured PMO enhances project success by providing governance, risk management, "
    "resource optimization, and process standardization. Investing in PMO activities leads to "
    "higher efficiency, better decision-making, and overall project excellence."
)

# Save the presentation
prs.save("PMO_Activities_Impact.pptx")

print("Presentation created successfully!")
